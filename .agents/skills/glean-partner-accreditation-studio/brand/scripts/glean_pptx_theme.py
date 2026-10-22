#!/usr/bin/env python3
"""
Glean-branded PPTX builder.

Encodes the measured geometry of Glean's seven master slide layouts (see
references/PRESENTATIONS.md) so decks come out on-brand without hand-placing shapes.

    from glean_pptx_theme import GleanDeck

    deck = GleanDeck()
    deck.cover("Partner-Led Services", "Operating model review - Q4")
    deck.section("Where we are today", "Three findings. One decision.")
    deck.three_column("What changed", [("Readiness", "Nine partners certified."), ...])
    deck.thank_you()
    deck.save("out.pptx")

Demo:
    python3 glean_pptx_theme.py --demo out.pptx

Requires: python-pptx. Background images are read from ../assets/backgrounds/ when present;
otherwise flat brand fills are used.
"""
from __future__ import annotations

import os
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- brand tokens
BLUE        = RGBColor(0x34, 0x3C, 0xED)
GREEN       = RGBColor(0xD8, 0xFD, 0x49)
OATMEAL     = RGBColor(0xF6, 0xF3, 0xEB)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
BLACK       = RGBColor(0x00, 0x00, 0x00)
NEAR_BLACK  = RGBColor(0x11, 0x13, 0x14)
BODY_GRAY   = RGBColor(0x7F, 0x7F, 0x7F)
MUTED_GRAY  = RGBColor(0x66, 0x66, 0x66)
MONO_GRAY   = RGBColor(0x72, 0x72, 0x72)
FOOTER_GRAY = RGBColor(0xB0, 0xAE, 0xA2)
PLACEHOLDER = RGBColor(0xED, 0xED, 0xED)
INDIGO      = RGBColor(0x40, 0x35, 0xC0)
ORCHID      = RGBColor(0xD0, 0x82, 0xC5)

FONT      = "PolySans"
FONT_ALT  = "DM Sans"          # Google Slides / Workspace fallback
FONT_MONO = "IBM Plex Mono"

# ---------------------------------------------------------------- geometry (inches)
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN      = 0.57             # content slides
COVER_INSET = 0.92             # cover / section slides
HEADER_TOP  = 0.64
BODY_TOP    = 3.16             # three-column body start
COL_PITCH   = 4.16
COL_WIDTH   = 3.87
FRAME_TOP   = 1.96
FRAME_H     = 5.12
FOOTER_TEXT = "Confidential | \u00a9 Glean Technologies, Inc."

_HERE = os.path.dirname(os.path.abspath(__file__))
_BG   = os.path.join(_HERE, "..", "assets", "backgrounds")
_LOGO = os.path.join(_HERE, "..", "assets", "logo")


def _asset(folder, name):
    p = os.path.normpath(os.path.join(folder, name))
    return p if os.path.exists(p) else None


class GleanDeck:
    """A 16:9 Glean-branded deck."""

    def __init__(self, use_alt_font: bool = False):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.font = FONT_ALT if use_alt_font else FONT

    # ---------------------------------------------------------- internals
    def _blank(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])

    def _bg(self, slide, style, fallback):
        img = _asset(_BG, f"glean-bg-{style}.png")
        if img:
            slide.shapes.add_picture(img, 0, 0,
                                     width=Inches(SLIDE_W), height=Inches(SLIDE_H))
        else:
            r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                       Inches(SLIDE_W), Inches(SLIDE_H))
            r.fill.solid()
            r.fill.fore_color.rgb = fallback
            r.line.fill.background()
            r.shadow.inherit = False

    def _text(self, slide, txt, left, top, width, height, size,
              *, bold=False, color=BLACK, align=PP_ALIGN.LEFT, mono=False,
              line=1.0, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                       Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, para_text in enumerate(str(txt).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line
            run = p.add_run()
            run.text = para_text
            f = run.font
            f.name = FONT_MONO if mono else self.font
            f.size = Pt(size)
            f.bold = bold
            f.color.rgb = color
        return box

    def _chrome(self, slide):
        """Footer + bottom-left avatar logo (content slides)."""
        self._text(slide, FOOTER_TEXT, MARGIN, 7.13, 12.70 - MARGIN, 0.22, 8,
                   color=FOOTER_GRAY, align=PP_ALIGN.RIGHT)
        logo = _asset(_LOGO, "glean-icon-electric-blue.png")
        if logo:
            slide.shapes.add_picture(logo, Inches(0.58), Inches(7.11), height=Inches(0.21))

    def _wordmark(self, slide, white=True):
        name = "glean-wordmark-white.png" if white else "glean-wordmark-electric-blue.png"
        logo = _asset(_LOGO, name)
        if logo:
            w = 1.32
            slide.shapes.add_picture(logo, Inches(12.70 - w), Inches(0.46), width=Inches(w))

    def _frame(self, slide, left, top, width, height, image=None):
        if image and os.path.exists(image):
            slide.shapes.add_picture(image, Inches(left), Inches(top),
                                     width=Inches(width), height=Inches(height))
            return
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                                     Inches(width), Inches(height))
        box.fill.solid()
        box.fill.fore_color.rgb = PLACEHOLDER
        box.line.fill.background()
        box.shadow.inherit = False
        try:                                    # subtle 8px-equivalent corner radius
            box.adjustments[0] = 0.02
        except Exception:
            pass

    # ---------------------------------------------------------- layouts
    def cover(self, title, subtitle=""):
        """Layout 1 - indigo cover. 96pt title, top-right wordmark."""
        s = self._blank()
        self._bg(s, "indigo", INDIGO)
        self._text(s, title, COVER_INSET, 2.05, 10.2, 2.9, 96,
                   bold=True, color=WHITE, line=1.0)
        if subtitle:
            self._text(s, subtitle, COVER_INSET, 5.24, 9.0, 0.6, 20,
                       bold=True, color=WHITE, line=1.2)
        self._wordmark(s)
        return s

    def section(self, title, subtitle=""):
        """Layout 2 - sunset divider. 88pt title, mono subtitle."""
        s = self._blank()
        self._bg(s, "sunset", ORCHID)
        self._text(s, title, 0.88, 2.34, 10.5, 2.6, 88, bold=True, color=WHITE, line=1.0)
        if subtitle:
            self._text(s, subtitle, 0.88, 5.10, 9.0, 0.6, 20,
                       color=WHITE, mono=True, line=1.2)
        self._wordmark(s)
        return s

    def three_column(self, header, columns):
        """Layout 3 - oatmeal, three text columns. columns: [(title, body), ...]"""
        s = self._blank()
        self._bg(s, "oatmeal", OATMEAL)
        self._text(s, header, MARGIN, HEADER_TOP, 12.20, 1.0, 36, bold=True, line=1.05)
        for i, col in enumerate(columns[:3]):
            left = MARGIN + i * COL_PITCH
            if isinstance(col, (tuple, list)):
                title, body = col[0], col[1]
                self._text(s, title, left, BODY_TOP, COL_WIDTH, 0.4, 16, bold=True)
                self._text(s, body, left, BODY_TOP + 0.42, COL_WIDTH, 2.4, 16,
                           color=BODY_GRAY, line=1.4)
            else:
                self._text(s, col, left, BODY_TOP, COL_WIDTH, 2.8, 16, line=1.4)
        self._chrome(s)
        return s

    def text_and_image(self, header, body, image=None):
        """Layout 4 - copy left, image frame right."""
        s = self._blank()
        self._bg(s, "oatmeal", OATMEAL)
        self._text(s, header, MARGIN, HEADER_TOP, 12.20, 1.0, 36, bold=True, line=1.05)
        self._text(s, body, MARGIN, 2.26, 4.60, 4.4, 16, color=BODY_GRAY, line=1.4)
        self._frame(s, 5.62, FRAME_TOP, 12.76 - 5.62, FRAME_H, image)
        self._chrome(s)
        return s

    def two_image(self, header, left_image=None, right_image=None):
        """Layout 5 - two equal image frames."""
        s = self._blank()
        self._bg(s, "oatmeal", OATMEAL)
        self._text(s, header, MARGIN, HEADER_TOP, 12.20, 1.0, 36, bold=True, line=1.05)
        self._frame(s, 0.47, FRAME_TOP, 6.12, FRAME_H, left_image)
        self._frame(s, 6.71, FRAME_TOP, 6.12, FRAME_H, right_image)
        self._chrome(s)
        return s

    def full_image(self, header, image=None, caption=""):
        """Layout 6 - one full-width image frame with optional caption."""
        s = self._blank()
        self._bg(s, "oatmeal", OATMEAL)
        self._text(s, header, MARGIN, HEADER_TOP, 12.20, 1.0, 36, bold=True, line=1.05)
        self._frame(s, 0.47, FRAME_TOP, 12.42, FRAME_H, image)
        if caption:
            self._text(s, caption, 0.72, FRAME_TOP + FRAME_H - 0.62, 11.0, 0.4, 16,
                       color=BLACK)
        self._chrome(s)
        return s

    def thank_you(self, title="Thank you", subtitle="", links=""):
        """Layout 7 - centered closing slide."""
        s = self._blank()
        self._bg(s, "sunset", ORCHID)
        self._text(s, title, 0, 2.87, SLIDE_W, 1.2, 72,
                   bold=True, color=WHITE, align=PP_ALIGN.CENTER, line=1.0)
        if subtitle:
            self._text(s, subtitle, 2.0, 4.15, SLIDE_W - 4.0, 0.9, 20,
                       color=WHITE, align=PP_ALIGN.CENTER, line=1.15)
        if links:
            self._text(s, links, 0, 6.65, SLIDE_W, 0.4, 16,
                       color=WHITE, align=PP_ALIGN.CENTER, mono=True)
        self._wordmark(s)
        return s

    # ---------------------------------------------------------- extras
    def stat_row(self, header, stats):
        """Oatmeal slide with 3-4 large stats. stats: [(number, label), ...]"""
        s = self._blank()
        self._bg(s, "oatmeal", OATMEAL)
        self._text(s, header, MARGIN, HEADER_TOP, 12.20, 1.0, 36, bold=True, line=1.05)
        n = max(len(stats), 1)
        width = 12.20 / n
        for i, (num, label) in enumerate(stats):
            left = MARGIN + i * width
            self._text(s, num, left, 3.10, width - 0.3, 1.0, 56, bold=True, line=1.0)
            self._text(s, label, left, 4.25, width - 0.3, 1.0, 16, color=BODY_GRAY, line=1.4)
        self._chrome(s)
        return s

    def save(self, path):
        self.prs.save(path)
        return path


def _demo(path):
    d = GleanDeck()
    d.cover("Glean Brand Style", "Every layout, one system")
    d.section("Foundations", "Color. Type. Space.")
    d.three_column(
        "Three things carry the brand",
        [("Oatmeal canvas", "Warm off-white, never pure white."),
         ("Electric Blue accent", "One point of emphasis per view."),
         ("Mono accessory", "Eyebrows, stats, and URLs in monospace.")],
    )
    d.text_and_image("Copy left, image right",
                     "Body copy sits in a 4.6in column so lines stay readable, "
                     "with the image frame filling the remaining width.")
    d.two_image("Paired comparison")
    d.full_image("One diagram, full width", caption="Caption sits inside the frame.")
    d.stat_row("By the numbers",
               [("28", "line icons"), ("7", "master layouts"), ("3", "brand surfaces")])
    d.thank_you("Thank you",
                "Let's build the super intelligent enterprise - together",
                "trust.glean.com | developers.glean.com")
    print("wrote", d.save(path))


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--demo", metavar="OUT.pptx")
    a = ap.parse_args()
    if a.demo:
        _demo(a.demo)
    else:
        ap.print_help()
